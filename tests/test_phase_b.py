"""
Phase B Tests — Schemas + SQLite + Workspace Scanner

Tests run against a temporary workspace directory.
No real DOCX parsing yet — we create dummy .docx files for scanning.
"""

import json
import os
import shutil
import sys
import tempfile
from pathlib import Path

# Add project root to path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from buildplan_knowledge import config, get_index_status, index_workspace, search_knowledge
from buildplan_knowledge.database import DatabaseManager
from buildplan_knowledge.schemas import (
    Chunk,
    ChunkMetadata,
    DocumentElement,
    DocumentSource,
    ElementType,
    FileIndexStatus,
    FileStatus,
    IndexWorkspaceResult,
    ParsedDocument,
    RetrievalItem,
    RetrievalResult,
)
from buildplan_knowledge.workspace import WorkspaceScanner


# ── Helpers ───────────────────────────────────────────────────────────────

def create_dummy_docx(path: Path, content: str = "test content") -> None:
    """Create a minimal valid .docx file."""
    # python-docx is available — use it to create real .docx files
    from docx import Document

    doc = Document()
    doc.add_heading("Test Document", level=1)
    doc.add_paragraph(content)
    doc.save(str(path))


def create_test_workspace(base_dir: Path) -> Path:
    """Create a test workspace with some .docx files and ignored dirs."""
    ws = base_dir / "test_workspace"
    ws.mkdir()

    # Create some .docx files
    create_dummy_docx(ws / "file1.docx", "Content of file 1")
    create_dummy_docx(ws / "file2.docx", "Content of file 2")

    # Create a subdirectory with a docx
    sub = ws / "subdir"
    sub.mkdir()
    create_dummy_docx(sub / "file3.docx", "Content of file 3")

    # Create ignored directories (should not be scanned)
    buildplan = ws / ".buildplan"
    buildplan.mkdir()
    create_dummy_docx(buildplan / "ignored.docx", "Should not be found")

    git_dir = ws / ".git"
    git_dir.mkdir()
    create_dummy_docx(git_dir / "ignored.docx", "Should not be found")

    pycache = ws / "__pycache__"
    pycache.mkdir()
    create_dummy_docx(pycache / "ignored.docx", "Should not be found")

    # Create non-docx files (should be ignored)
    (ws / "readme.txt").write_text("not a docx")
    (ws / "data.csv").write_text("a,b,c")

    return ws


# ── Test functions ────────────────────────────────────────────────────────

def test_workspace_scanner():
    """Test that WorkspaceScanner discovers the right files."""
    print("=" * 60)
    print("TEST: Workspace Scanner")
    print("=" * 60)

    with tempfile.TemporaryDirectory() as tmpdir:
        ws = create_test_workspace(Path(tmpdir))
        scanner = WorkspaceScanner(ws)
        discovered = scanner.scan()

        # Should find exactly 3 .docx files (file1, file2, subdir/file3)
        assert len(discovered) == 3, f"Expected 3 files, got {len(discovered)}"

        paths = {f.relative_path for f in discovered}
        assert "file1.docx" in paths
        assert "file2.docx" in paths
        assert "subdir/file3.docx" in paths

        # Check that each file has a hash and document_id
        for f in discovered:
            assert len(f.file_hash) == 64, f"Expected SHA-256 hash, got {len(f.file_hash)}"
            assert f.document_id, "document_id should not be empty"
            assert f.file_size > 0, "file_size should be positive"
            assert f.file_type == ".docx"

        print(f"  ✓ Discovered {len(discovered)} files")
        for f in discovered:
            print(f"    - {f.relative_path} ({f.file_size} bytes, hash={f.file_hash[:8]}...)")


def test_database_operations():
    """Test SQLite database operations."""
    print("\n" + "=" * 60)
    print("TEST: Database Operations")
    print("=" * 60)

    with tempfile.TemporaryDirectory() as tmpdir:
        ws = Path(tmpdir) / "test_ws"
        ws.mkdir()

        with DatabaseManager(ws) as db:
            # Test workspace ID creation
            ws_id = db.get_or_create_workspace_id()
            assert ws_id, "workspace_id should not be empty"
            assert len(ws_id) == 36, f"Expected UUID format, got: {ws_id}"
            print(f"  ✓ workspace_id created: {ws_id}")

            # Test workspace ID persistence
            ws_id2 = db.get_or_create_workspace_id()
            assert ws_id == ws_id2, "workspace_id should be stable across calls"
            print(f"  ✓ workspace_id is stable: {ws_id2}")

            # Verify workspace.json was created
            meta_path = config.get_workspace_meta_path(ws)
            assert meta_path.exists(), "workspace.json should exist"
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
            assert meta["workspace_id"] == ws_id
            print(f"  ✓ workspace.json persisted correctly")

            # Test file upsert
            db.upsert_file(
                document_id="doc-001",
                relative_path="test.docx",
                file_name="test.docx",
                file_type=".docx",
                file_size=1024,
                file_hash="abc123",
                modified_time=1000.0,
            )
            f = db.get_file_by_path("test.docx")
            assert f is not None
            assert f["document_id"] == "doc-001"
            assert f["indexed"] == 0
            print(f"  ✓ File upsert works")

            # Test mark_indexed
            db.mark_indexed("test.docx", chunk_count=5)
            f = db.get_file_by_path("test.docx")
            assert f["indexed"] == 1
            assert f["chunk_count"] == 5
            assert f["indexed_at"] is not None
            print(f"  ✓ mark_indexed works")

            # Test mark_skipped
            db.mark_skipped("test.docx")
            f = db.get_file_by_path("test.docx")
            assert f["indexed"] == 1
            print(f"  ✓ mark_skipped works")

            # Test mark_failed
            db.mark_failed("test.docx", "parse error")
            f = db.get_file_by_path("test.docx")
            assert f["indexed"] == 0
            assert f["error_message"] == "parse error"
            print(f"  ✓ mark_failed works")

            # Test get_stats
            db.mark_indexed("test.docx", chunk_count=5)
            stats = db.get_stats()
            assert stats["total_files"] == 1
            assert stats["indexed_files"] == 1
            print(f"  ✓ get_stats works: {stats}")

            # Test delete_file
            db.delete_file("test.docx")
            assert db.get_file_by_path("test.docx") is None
            print(f"  ✓ delete_file works")


def test_index_workspace_api():
    """Test the index_workspace public API."""
    print("\n" + "=" * 60)
    print("TEST: index_workspace API")
    print("=" * 60)

    with tempfile.TemporaryDirectory() as tmpdir:
        ws = create_test_workspace(Path(tmpdir))

        # First index
        result = index_workspace(str(ws))
        assert isinstance(result, IndexWorkspaceResult)
        assert result.discovered == 3
        assert result.indexed == 3  # all new
        assert result.skipped == 0
        assert result.failed == 0
        assert result.workspace_id, "workspace_id should be set"
        print(f"  ✓ First index: discovered={result.discovered}, indexed={result.indexed}")

        # Second index (unchanged files → all skipped)
        result2 = index_workspace(str(ws))
        assert result2.discovered == 3
        assert result2.indexed == 0
        assert result2.skipped == 3
        print(f"  ✓ Second index (unchanged): skipped={result2.skipped}")

        # Verify workspace_id is stable across calls
        assert result.workspace_id == result2.workspace_id
        print(f"  ✓ workspace_id stable across index calls")

        # Modify one file and re-index
        create_dummy_docx(ws / "file1.docx", "MODIFIED content")
        result3 = index_workspace(str(ws))
        assert result3.discovered == 3
        assert result3.indexed == 1  # file1 changed
        assert result3.skipped == 2  # file2, file3 unchanged
        print(f"  ✓ After modify: indexed={result3.indexed}, skipped={result3.skipped}")

        # Delete one file and re-index
        (ws / "subdir" / "file3.docx").unlink()
        result4 = index_workspace(str(ws))
        assert result4.discovered == 2
        # file3 should be removed from DB
        status = get_index_status(str(ws))
        assert status["discovered_files"] == 2
        print(f"  ✓ After delete: discovered={result4.discovered}, DB records={status['discovered_files']}")


def test_get_index_status_api():
    """Test the get_index_status public API."""
    print("\n" + "=" * 60)
    print("TEST: get_index_status API")
    print("=" * 60)

    with tempfile.TemporaryDirectory() as tmpdir:
        ws = create_test_workspace(Path(tmpdir))

        # Index first
        index_workspace(str(ws))

        # Get status
        status = get_index_status(str(ws))
        assert status["workspace_id"]
        assert status["discovered_files"] == 3
        assert status["indexed_files"] == 3
        assert status["failed_files"] == 0
        assert len(status["files"]) == 3

        for f in status["files"]:
            assert f["status"] == "indexed"
            assert f["relative_path"]

        print(f"  ✓ Status: {status['discovered_files']} files, {status['total_chunks']} chunks")
        for f in status["files"]:
            print(f"    - {f['relative_path']}: {f['status']}")


def test_search_knowledge_stub():
    """Test that search_knowledge returns empty result in Phase B."""
    print("\n" + "=" * 60)
    print("TEST: search_knowledge stub")
    print("=" * 60)

    with tempfile.TemporaryDirectory() as tmpdir:
        ws = create_test_workspace(Path(tmpdir))
        index_workspace(str(ws))

        result = search_knowledge(str(ws), "test query")
        assert isinstance(result, RetrievalResult)
        assert result.query == "test query"
        assert result.scope == "workspace"
        assert result.results == []
        print(f"  ✓ search_knowledge returns empty result (stub)")


def test_schema_dataclasses():
    """Test that all schema dataclasses instantiate correctly."""
    print("\n" + "=" * 60)
    print("TEST: Schema Dataclasses")
    print("=" * 60)

    # DocumentElement
    elem = DocumentElement(
        element_type=ElementType.HEADING,
        text="4 施工部署",
        level=1,
        heading_path=["4 施工部署"],
    )
    assert elem.element_type == ElementType.HEADING
    assert elem.level == 1
    print(f"  ✓ DocumentElement OK")

    # ParsedDocument
    parsed = ParsedDocument(
        document_id="doc-1",
        file_name="test.docx",
        relative_path="test.docx",
        file_type=".docx",
        elements=[elem],
    )
    assert len(parsed.elements) == 1
    print(f"  ✓ ParsedDocument OK")

    # Chunk
    chunk = Chunk(
        chunk_id="chunk-001",
        document_id="doc-1",
        chunk_index=0,
        text="原始正文",
        embedding_text="章节：4 施工部署\n\n原始正文",
        heading_path=["4 施工部署"],
        metadata=ChunkMetadata(
            workspace_id="ws-1",
            document_id="doc-1",
            file_name="test.docx",
            relative_path="test.docx",
            file_type=".docx",
            heading_path=["4 施工部署"],
        ),
    )
    assert chunk.text == "原始正文"
    assert "4 施工部署" in chunk.embedding_text
    print(f"  ✓ Chunk OK")

    # RetrievalResult
    ret = RetrievalResult(
        query="test",
        scope="workspace",
        results=[
            RetrievalItem(
                chunk_id="c1",
                text="answer",
                score=0.95,
                source=DocumentSource(
                    document_id="d1",
                    file_name="f.docx",
                    relative_path="f.docx",
                    heading_path=["Section 1"],
                ),
            )
        ],
    )
    assert len(ret.results) == 1
    assert ret.results[0].score == 0.95
    print(f"  ✓ RetrievalResult OK")

    # FileStatus
    fs = FileStatus(
        relative_path="test.docx",
        file_name="test.docx",
        status=FileIndexStatus.INDEXED,
        chunk_count=10,
    )
    assert fs.status == FileIndexStatus.INDEXED
    print(f"  ✓ FileStatus OK")


def test_config_values():
    """Test that config values are reasonable."""
    print("\n" + "=" * 60)
    print("TEST: Config Values")
    print("=" * 60)

    assert config.CHUNK_TARGET_TOKENS == 600
    assert config.CHUNK_OVERLAP_TOKENS == 100
    assert config.CHARS_PER_TOKEN == 1.5
    assert config.CHUNK_TARGET_CHARS == 900
    assert config.CHUNK_OVERLAP_CHARS == 150
    assert config.TOP_K_DEFAULT == 5
    assert ".docx" in config.SUPPORTED_EXTENSIONS
    assert ".buildplan" in config.IGNORED_DIRS
    assert ".git" in config.IGNORED_DIRS
    print(f"  ✓ All config values correct")
    print(f"    CHUNK_TARGET_TOKENS={config.CHUNK_TARGET_TOKENS}")
    print(f"    CHUNK_OVERLAP_TOKENS={config.CHUNK_OVERLAP_TOKENS}")
    print(f"    CHARS_PER_TOKEN={config.CHARS_PER_TOKEN}")
    print(f"    SUPPORTED_EXTENSIONS={config.SUPPORTED_EXTENSIONS}")


def test_buildplan_dir_structure():
    """Test that .buildplan/ directory is created correctly."""
    print("\n" + "=" * 60)
    print("TEST: .buildplan/ Directory Structure")
    print("=" * 60)

    with tempfile.TemporaryDirectory() as tmpdir:
        ws = create_test_workspace(Path(tmpdir))

        index_workspace(str(ws))

        buildplan_dir = config.get_buildplan_dir(ws)
        assert buildplan_dir.exists(), ".buildplan/ should exist"
        assert buildplan_dir.is_dir()

        db_path = config.get_db_path(ws)
        assert db_path.exists(), "project.db should exist"

        meta_path = config.get_workspace_meta_path(ws)
        assert meta_path.exists(), "workspace.json should exist"

        # Verify workspace.json content
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
        assert "workspace_id" in meta

        # Verify DB has the files table
        import sqlite3
        conn = sqlite3.connect(str(db_path))
        tables = [r[0] for r in conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table'"
        ).fetchall()]
        assert "files" in tables
        conn.close()

        print(f"  ✓ .buildplan/ created at: {buildplan_dir}")
        print(f"  ✓ project.db at: {db_path}")
        print(f"  ✓ workspace.json at: {meta_path}")
        print(f"  ✓ Tables in DB: {tables}")


# ── Main ──────────────────────────────────────────────────────────────────

def run_all_tests():
    """Run all Phase B tests."""
    print("\n" + "█" * 60)
    print("  PHASE B TESTS — Schemas + SQLite + Workspace Scanner")
    print("█" * 60)

    tests = [
        test_config_values,
        test_schema_dataclasses,
        test_workspace_scanner,
        test_database_operations,
        test_buildplan_dir_structure,
        test_index_workspace_api,
        test_get_index_status_api,
        test_search_knowledge_stub,
    ]

    passed = 0
    failed = 0

    for test_fn in tests:
        try:
            test_fn()
            passed += 1
        except Exception as e:
            failed += 1
            print(f"\n  ✗ FAILED: {test_fn.__name__}")
            print(f"    Error: {e}")
            import traceback
            traceback.print_exc()

    print("\n" + "█" * 60)
    print(f"  RESULTS: {passed} passed, {failed} failed, {passed + failed} total")
    print("█" * 60)

    if failed > 0:
        sys.exit(1)
    else:
        print("\n  ✅ All Phase B tests passed!")


if __name__ == "__main__":
    run_all_tests()
